"""Lecture Excel/CSV et contrôles avant import."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook

from app.services.ecoles import annee_academique_courante, ecole_officielle

_NOM = re.compile(r"^[A-Za-zÀ-ÖØ-öø-ÿ'’\- ]{2,80}$")
_FILIERE = re.compile(r"^[A-Za-zÀ-ÖØ-öø-ÿ0-9'’\-/&, ]{2,150}$")
_MATRICULE = re.compile(r"^[A-Z0-9][A-Z0-9-]{2,30}$")
_ANNEE = re.compile(r"^(\d{4})-(\d{4})$")

_EMAIL = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,24}$")

_COLONNES = {
    "prenom": "prenom",
    "nom": "nom",
    "matricule": "matricule",
    "filiere": "filiere",
    "ecole": "filiere",
    "annee": "annee_academique",
    "annee_academique": "annee_academique",
    "date_validite": "date_validite",
    "date_de_validite": "date_validite",
    "photo": "photo",
    "email": "email",
    "mail": "email",
    "courriel": "email",
    "email_professionnel": "email",
    "sexe": "sexe",
    "genre": "sexe",
}


def lire_sexe(valeur: str | None, obligatoire: bool = False) -> str:
    """M ou F. Le sexe ne s'imprime pas : il choisit seulement étudiant ou étudiante."""
    brut = unicodedata.normalize("NFD", (valeur or "").strip().lower())
    brut = "".join(car for car in brut if unicodedata.category(car) != "Mn")
    if brut in {"m", "masculin", "homme"}:
        return "M"
    if brut in {"f", "feminin", "femme"}:
        return "F"
    if not brut and not obligatoire:
        return ""
    raise ValueError("Choisissez Masculin ou Féminin.")


@dataclass
class Fiche:
    prenom: str
    nom: str
    matricule: str
    filiere: str
    annee_academique: str
    date_validite: date
    email: str = ""
    avertissement: str | None = None
    sexe: str = ""


@dataclass
class LigneFichier:
    numero: int
    valeurs: dict[str, str]


def _sans_accent(texte: str) -> str:
    normalise = unicodedata.normalize("NFD", texte.strip().lower())
    sans = "".join(c for c in normalise if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", sans).strip("_")


def _propre(texte: str) -> str:
    return re.sub(r"\s+", " ", (texte or "")).strip()


def validite_par_defaut(annee: str) -> date:
    """Trois ans à compter du 1er septembre. 2026-2027 est valable jusqu'au 31/08/2029."""
    debut = int(str(annee).split("-")[0])
    return date(debut + 3, 8, 31)


def _lire_date(texte: str) -> date | None:
    texte = texte.strip()
    if not texte:
        return None
    if isinstance(texte, datetime):
        return texte.date()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(texte), fmt).date()
        except ValueError:
            continue
    return None


def controler_email(email: str | None) -> str:
    email = (email or "").strip().lower()
    if not email:
        return ""
    if len(email) > 120 or not _EMAIL.match(email):
        raise ValueError("l'e-mail professionnel n'est pas valide.")
    return email


def controler_fiche(
    prenom: str,
    nom: str,
    matricule: str,
    filiere: str,
    annee: str,
    date_texte: str | None = None,
    email: str | None = None,
    sexe: str | None = None,
    sexe_obligatoire: bool = False,
) -> Fiche:
    prenom = _propre(prenom)
    nom = _propre(nom)
    matricule = re.sub(r"\s+", "", matricule or "").upper()
    annee = annee_academique_courante()

    if not _NOM.match(prenom):
        raise ValueError("le prénom est vide ou contient un caractère non autorisé.")
    if not _NOM.match(nom):
        raise ValueError("le nom est vide ou contient un caractère non autorisé.")
    if not _MATRICULE.match(matricule):
        raise ValueError("le matricule doit ressembler à AIIM-2026-0142 (lettres, chiffres, tirets).")
    try:
        filiere = ecole_officielle(filiere)
    except ValueError as exc:
        raise ValueError(str(exc)) from exc

    avertissement = None
    date_validite = validite_par_defaut(annee)

    return Fiche(
        prenom,
        nom,
        matricule,
        filiere,
        annee,
        date_validite,
        controler_email(email),
        avertissement,
        lire_sexe(sexe, obligatoire=sexe_obligatoire),
    )


def _cellule(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, date):
        return valeur.strftime("%d/%m/%Y")
    return str(valeur).strip()


def lire_tableau(contenu: bytes, nom_fichier: str) -> list[LigneFichier]:
    if len(contenu) > 2_000_000:
        raise ValueError("Le fichier dépasse 2 Mo.")
    nom = nom_fichier.lower()
    if nom.endswith(".xlsx"):
        return _lire_xlsx(contenu)
    if nom.endswith(".csv") or nom.endswith(".txt"):
        return _lire_csv(contenu)
    raise ValueError("Format accepté : .csv ou .xlsx (enregistrez l'ancien .xls en .xlsx).")


def _mapper_entetes(entetes: list[str]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, entete in enumerate(entetes):
        cle = _COLONNES.get(_sans_accent(entete))
        if cle and cle not in mapping.values():
            mapping[index] = cle
    manquantes = [c for c in ("prenom", "nom", "matricule", "filiere") if c not in mapping.values()]
    if manquantes:
        raise ValueError(
            "Colonnes manquantes : "
            + ", ".join(manquantes)
            + ". La première ligne doit contenir prenom, nom, matricule, ecole."
        )
    return mapping


def _lignes_depuis_rangees(rangees: list[list]) -> list[LigneFichier]:
    if not rangees:
        raise ValueError("Le fichier est vide.")
    mapping = _mapper_entetes([_cellule(c) for c in rangees[0]])
    lignes: list[LigneFichier] = []
    for numero, rangee in enumerate(rangees[1:], start=2):
        valeurs = {champ: "" for champ in set(mapping.values())}
        for index, champ in mapping.items():
            if index < len(rangee):
                valeurs[champ] = _cellule(rangee[index])
        if not any(valeurs.values()):
            continue
        lignes.append(LigneFichier(numero, valeurs))
    if len(lignes) > 5000:
        raise ValueError("Maximum 5000 lignes par import.")
    return lignes


def _lire_csv(contenu: bytes) -> list[LigneFichier]:
    texte = None
    for encodage in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texte = contenu.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    if texte is None:
        raise ValueError("Impossible de lire le fichier. Enregistrez-le en UTF-8.")
    echantillon = texte[:2000]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
        dialecte.delimiter = ";" if echantillon.count(";") >= echantillon.count(",") else ","
    lecteur = csv.reader(io.StringIO(texte), dialecte)
    return _lignes_depuis_rangees(list(lecteur))


def _lire_xlsx(contenu: bytes) -> list[LigneFichier]:
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Fichier Excel illisible.") from exc
    feuille = classeur.active
    rangees = [list(ligne) for ligne in feuille.iter_rows(values_only=True)]
    classeur.close()
    return _lignes_depuis_rangees(rangees)


def analyser_lignes(
    lignes: list[LigneFichier], matricules_connus: set[str]
) -> tuple[list[tuple[LigneFichier, Fiche]], list[str], list[str]]:
    pretes: list[tuple[LigneFichier, Fiche]] = []
    erreurs: list[str] = []
    avertissements: list[str] = []
    vus: set[str] = set()

    for ligne in lignes:
        valeurs = ligne.valeurs
        try:
            fiche = controler_fiche(
                valeurs.get("prenom", ""),
                valeurs.get("nom", ""),
                valeurs.get("matricule", ""),
                valeurs.get("filiere", ""),
                valeurs.get("annee_academique", ""),
                valeurs.get("date_validite") or None,
                valeurs.get("email") or None,
                valeurs.get("sexe"),
            )
        except ValueError as exc:
            erreurs.append(f"Ligne {ligne.numero} : {exc}")
            continue
        if fiche.matricule in vus:
            erreurs.append(f"Ligne {ligne.numero} : le matricule {fiche.matricule} est en double dans le fichier.")
            continue
        if fiche.matricule in matricules_connus:
            erreurs.append(f"Ligne {ligne.numero} : le matricule {fiche.matricule} existe déjà.")
            continue
        vus.add(fiche.matricule)
        if fiche.avertissement:
            avertissements.append(f"Ligne {ligne.numero} ({fiche.matricule}) : {fiche.avertissement}.")
        pretes.append((ligne, fiche))
    return pretes, erreurs, avertissements
